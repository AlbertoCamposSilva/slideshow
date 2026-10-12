"""Gerenciamento de curtidas (Likes/Favoritos), unlikes e configurações persistentes em SQLite.

O banco é a fonte da verdade. A planilha Excel (.xlsx) passa a ser apenas formato de
intercâmbio: `export_xlsx` gera uma cópia legível e `import_xlsx` substitui todo o conteúdo
do banco pelo da planilha (replace all).
"""

import os
import sqlite3
from contextlib import closing
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Set, Tuple

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill


DEFAULT_DB_FILE = Path.home() / "slideshow_favoritos.db"

KIND_LIKE = "like"
KIND_UNLIKE = "unlike"

_SHEET_BY_KIND = {KIND_LIKE: "Favoritos", KIND_UNLIKE: "Unlikes"}
_COLOR_BY_KIND = {KIND_LIKE: "2F5597", KIND_UNLIKE: "C00000"}
_PHOTO_HEADERS = [
    "Caminho Completo",
    "Nome do Arquivo",
    "Data/Hora",
    "Resolução",
    "Tamanho (KB)",
    "Pasta de Origem",
]
_CONFIG_SHEET = "Configuracoes"
_CONFIG_HEADERS = ["Chave", "Valor", "Ultima_Atualizacao"]

_SCHEMA = """
CREATE TABLE IF NOT EXISTS photos (
    path_norm  TEXT PRIMARY KEY,
    kind       TEXT NOT NULL CHECK (kind IN ('like', 'unlike')),
    path       TEXT NOT NULL,
    filename   TEXT NOT NULL,
    created_at TEXT NOT NULL,
    resolution TEXT NOT NULL DEFAULT '',
    size_kb    REAL NOT NULL DEFAULT 0,
    folder     TEXT NOT NULL DEFAULT ''
);
CREATE INDEX IF NOT EXISTS idx_photos_kind ON photos(kind);
CREATE TABLE IF NOT EXISTS config (
    key        TEXT PRIMARY KEY,
    value      TEXT NOT NULL,
    updated_at TEXT NOT NULL
);
"""

# (path_norm, kind, path, filename, created_at, resolution, size_kb, folder)
_PhotoRow = Tuple[str, str, str, str, str, str, float, str]
# (key, value, updated_at)
_ConfigRow = Tuple[str, str, str]


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


class FavoritesManager:
    """Gerencia favoritos, unlikes e configurações persistentes em um banco SQLite."""

    def __init__(self, db_path: Optional[Path] = None):
        self.db_path = Path(db_path) if db_path else DEFAULT_DB_FILE
        self._favorites: Set[str] = set()
        self._unlikes: Set[str] = set()
        self._init_db()
        self._load()

    # --- Infraestrutura ---

    def _connect(self) -> sqlite3.Connection:
        return sqlite3.connect(self.db_path, timeout=10)

    def _init_db(self):
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        with closing(self._connect()) as conn, conn:
            conn.executescript(_SCHEMA)

    @staticmethod
    def _normalize_path(path: str) -> str:
        """Normaliza caminho para comparação exata (insensível a maiúsculas no Windows)."""
        return os.path.normcase(os.path.abspath(path))

    def _load(self):
        """Recarrega os conjuntos em memória a partir do banco."""
        self._favorites.clear()
        self._unlikes.clear()
        with closing(self._connect()) as conn:
            for path_norm, kind in conn.execute("SELECT path_norm, kind FROM photos"):
                (self._favorites if kind == KIND_LIKE else self._unlikes).add(path_norm)

    # --- Likes e Unlikes ---

    def is_favorite(self, path: str) -> bool:
        """Verifica se o arquivo está favoritado."""
        return self._normalize_path(path) in self._favorites

    def is_unliked(self, path: str) -> bool:
        """Verifica se o arquivo está marcado com unlike."""
        return self._normalize_path(path) in self._unlikes

    def add_favorite(self, path: str, metadata: Optional[Dict[str, Any]] = None) -> bool:
        """Adiciona aos favoritos (substitui um unlike existente)."""
        return self._set_kind(path, KIND_LIKE, metadata)

    def add_unlike(self, path: str, metadata: Optional[Dict[str, Any]] = None) -> bool:
        """Adiciona aos unlikes (substitui um favorito existente)."""
        return self._set_kind(path, KIND_UNLIKE, metadata)

    def remove_favorite(self, path: str) -> bool:
        """Remove a imagem dos favoritos."""
        return self._remove(path, KIND_LIKE)

    def remove_unlike(self, path: str) -> bool:
        """Remove a imagem dos unlikes."""
        return self._remove(path, KIND_UNLIKE)

    def _set_kind(self, path: str, kind: str, metadata: Optional[Dict[str, Any]]) -> bool:
        norm = self._normalize_path(path)
        mine, other = (
            (self._favorites, self._unlikes) if kind == KIND_LIKE else (self._unlikes, self._favorites)
        )
        if norm in mine:
            return False

        resolution = ""
        size_kb = 0.0
        if metadata:
            resolution = str(metadata.get("resolution", ""))
            size_kb = float(metadata.get("size_kb", 0.0))
        else:
            try:
                size_kb = round(os.path.getsize(path) / 1024, 2)
            except OSError:
                pass

        try:
            with closing(self._connect()) as conn, conn:
                # PRIMARY KEY em path_norm garante exclusão mútua entre like e unlike.
                conn.execute(
                    "INSERT OR REPLACE INTO photos VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
                    (norm, kind, path, os.path.basename(path), _now(),
                     resolution, size_kb, os.path.dirname(path)),
                )
        except sqlite3.Error as e:
            print(f"[Favoritos] Erro ao gravar '{path}': {e}")
            return False

        other.discard(norm)
        mine.add(norm)
        return True

    def _remove(self, path: str, kind: str) -> bool:
        norm = self._normalize_path(path)
        bucket = self._favorites if kind == KIND_LIKE else self._unlikes
        if norm not in bucket:
            return False
        try:
            with closing(self._connect()) as conn, conn:
                conn.execute("DELETE FROM photos WHERE path_norm = ? AND kind = ?", (norm, kind))
        except sqlite3.Error as e:
            print(f"[Favoritos] Erro ao remover '{path}': {e}")
            return False
        bucket.discard(norm)
        return True

    # --- Persistência de Configurações ---

    def load_config(self) -> Dict[str, Any]:
        """Carrega o dicionário de configurações salvas."""
        try:
            with closing(self._connect()) as conn:
                return dict(conn.execute("SELECT key, value FROM config"))
        except sqlite3.Error as e:
            print(f"[Config] Erro ao carregar configurações de '{self.db_path}': {e}")
            return {}

    def save_config(self, config_dict: Dict[str, Any]):
        """Salva ou atualiza pares chave-valor de configuração."""
        now = _now()
        rows = [(str(k).strip(), str(v), now) for k, v in config_dict.items()]
        try:
            with closing(self._connect()) as conn, conn:
                conn.executemany("INSERT OR REPLACE INTO config VALUES (?, ?, ?)", rows)
        except sqlite3.Error as e:
            print(f"[Config] Erro ao salvar configurações em '{self.db_path}': {e}")

    # --- Exportação / Importação (Excel) ---

    def export_xlsx(self, xlsx_path: Path) -> Tuple[int, int]:
        """Exporta favoritos, unlikes e configurações para um .xlsx. Retorna (likes, unlikes)."""
        xlsx_path = Path(xlsx_path)
        wb = openpyxl.Workbook()
        wb.remove(wb.active)

        with closing(self._connect()) as conn:
            counts = {}
            for kind in (KIND_LIKE, KIND_UNLIKE):
                ws = wb.create_sheet(_SHEET_BY_KIND[kind])
                self._style_header(ws, _PHOTO_HEADERS, _COLOR_BY_KIND[kind])
                rows = conn.execute(
                    "SELECT path, filename, created_at, resolution, size_kb, folder "
                    "FROM photos WHERE kind = ? ORDER BY created_at, rowid",
                    (kind,),
                ).fetchall()
                for row in rows:
                    ws.append(list(row))
                for letter, width in zip("ABCDEF", (60, 30, 20, 15, 15, 45)):
                    ws.column_dimensions[letter].width = width
                counts[kind] = len(rows)

            ws_cfg = wb.create_sheet(_CONFIG_SHEET)
            self._style_header(ws_cfg, _CONFIG_HEADERS, "375623")
            for row in conn.execute("SELECT key, value, updated_at FROM config ORDER BY key"):
                ws_cfg.append(list(row))
            for letter, width in zip("ABC", (25, 30, 22)):
                ws_cfg.column_dimensions[letter].width = width

        xlsx_path.parent.mkdir(parents=True, exist_ok=True)
        tmp = xlsx_path.with_name(xlsx_path.name + ".tmp")
        wb.save(tmp)
        os.replace(tmp, xlsx_path)
        return counts[KIND_LIKE], counts[KIND_UNLIKE]

    def import_xlsx(self, xlsx_path: Path) -> Tuple[int, int]:
        """Substitui TODO o conteúdo do banco pelo da planilha. Retorna (likes, unlikes).

        A planilha é lida e validada por completo antes de qualquer alteração; em caso de erro
        o banco permanece intacto. A gravação ocorre numa única transação.
        """
        xlsx_path = Path(xlsx_path)
        wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
        try:
            known = {_SHEET_BY_KIND[KIND_LIKE], _SHEET_BY_KIND[KIND_UNLIKE], _CONFIG_SHEET}
            if not known & set(wb.sheetnames):
                raise ValueError(
                    "Planilha inválida: nenhuma das abas 'Favoritos', 'Unlikes' ou "
                    "'Configuracoes' foi encontrada."
                )

            photos: Dict[str, _PhotoRow] = {}
            # Unlikes primeiro: se o mesmo caminho estiver nas duas abas, o Like prevalece.
            for kind in (KIND_UNLIKE, KIND_LIKE):
                sheet = _SHEET_BY_KIND[kind]
                if sheet not in wb.sheetnames:
                    continue
                for row in wb[sheet].iter_rows(min_row=2, values_only=True):
                    if not row or not row[0]:
                        continue
                    row = list(row) + [None] * (6 - len(row))
                    raw = str(row[0])
                    norm = self._normalize_path(raw)
                    try:
                        size_kb = float(row[4]) if row[4] not in (None, "") else 0.0
                    except (TypeError, ValueError):
                        size_kb = 0.0
                    photos[norm] = (
                        norm, kind, raw,
                        str(row[1]) if row[1] else os.path.basename(raw),
                        str(row[2]) if row[2] else _now(),
                        str(row[3]) if row[3] else "",
                        size_kb,
                        str(row[5]) if row[5] else os.path.dirname(raw),
                    )

            config: List[_ConfigRow] = []
            if _CONFIG_SHEET in wb.sheetnames:
                for row in wb[_CONFIG_SHEET].iter_rows(min_row=2, values_only=True):
                    if row and row[0] is not None and row[1] is not None:
                        config.append((
                            str(row[0]).strip(),
                            str(row[1]),
                            str(row[2]) if len(row) > 2 and row[2] else _now(),
                        ))
        finally:
            wb.close()

        with closing(self._connect()) as conn, conn:
            conn.execute("DELETE FROM photos")
            conn.execute("DELETE FROM config")
            conn.executemany("INSERT INTO photos VALUES (?, ?, ?, ?, ?, ?, ?, ?)", photos.values())
            conn.executemany("INSERT OR REPLACE INTO config VALUES (?, ?, ?)", config)

        self._load()
        return self.favorites_count, self.unlikes_count

    @staticmethod
    def _style_header(ws, headers: List[str], color: str):
        ws.append(headers)
        font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill = PatternFill(start_color=color, end_color=color, fill_type="solid")
        align = Alignment(horizontal="center", vertical="center")
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font, cell.fill, cell.alignment = font, fill, align

    @property
    def favorites_count(self) -> int:
        """Total de fotos favoritadas."""
        return len(self._favorites)

    @property
    def unlikes_count(self) -> int:
        """Total de fotos unliked."""
        return len(self._unlikes)
